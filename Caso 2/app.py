from flask import Flask, request, render_template_string, send_file
from openpyxl import Workbook, load_workbook
import os

app = Flask(__name__)

ARCHIVO = "miembros_mesa.xlsx"

HTML = """
<!DOCTYPE html>
<html lang="es">

<head>

    <meta charset="UTF-8">

    <title>Miembro de Mesa</title>

    <style>

        body {
            font-family: Arial, sans-serif;
            background: #f2f2f2;
            padding: 40px;
        }

        .container {
            background: white;
            max-width: 700px;
            margin: auto;
            padding: 30px;
            border-radius: 12px;
            box-shadow: 0 0 10px #ccc;
        }

        h1 {
            text-align: center;
        }

        label {
            display: block;
            margin-top: 15px;
            font-weight: bold;
        }

        input {
            width: 95%;
            padding: 12px;
            margin-top: 5px;
            border: 1px solid #ccc;
            border-radius: 5px;
        }

        button {
            margin-top: 20px;
            padding: 12px 25px;
            background: #333;
            color: white;
            border: none;
            border-radius: 5px;
            cursor: pointer;
        }

        button:hover {
            background: #555;
        }

        .mensaje {
            margin-top: 20px;
            text-align: center;
            font-weight: bold;
        }

    </style>

</head>

<body>

<div class="container">

    <h1>🗳️ Registro de Miembro de Mesa</h1>

    <form method="POST">

        <label>DNI</label>

        <input
            type="text"
            name="dni"
            maxlength="8"
            placeholder="Ingrese DNI"
            required
        >

        <label>Región</label>

        <input
            type="text"
            name="region"
            placeholder="Ejemplo: Lima"
            required
        >

        <label>Provincia</label>

        <input
            type="text"
            name="provincia"
            placeholder="Ingrese provincia"
            required
        >

        <label>Distrito</label>

        <input
            type="text"
            name="distrito"
            placeholder="Ingrese distrito"
            required
        >

        <label>Dirección del local de votación</label>

        <input
            type="text"
            name="direccion"
            placeholder="Ingrese dirección"
            required
        >

        <button type="submit">
            Registrar información
        </button>

    </form>

    {% if mensaje %}

        <div class="mensaje">
            {{ mensaje }}
        </div>

    {% endif %}

</div>

</body>

</html>
"""


def crear_excel():

    if not os.path.exists(ARCHIVO):

        wb = Workbook()

        ws = wb.active

        ws.title = "Miembros de Mesa"

        ws.append([
            "DNI",
            "Region",
            "Provincia",
            "Distrito",
            "Direccion del local de votacion"
        ])

        wb.save(ARCHIVO)


def guardar_datos(
    dni,
    region,
    provincia,
    distrito,
    direccion
):

    crear_excel()

    wb = load_workbook(ARCHIVO)

    ws = wb["Miembros de Mesa"]

    ws.append([
        dni,
        region,
        provincia,
        distrito,
        direccion
    ])

    wb.save(ARCHIVO)


@app.route("/", methods=["GET", "POST"])
def index():

    mensaje = ""

    if request.method == "POST":

        dni = request.form["dni"]
        region = request.form["region"]
        provincia = request.form["provincia"]
        distrito = request.form["distrito"]
        direccion = request.form["direccion"]

        if not dni.isdigit():

            mensaje = "El DNI solo debe contener números."

        elif len(dni) != 8:

            mensaje = "El DNI debe tener 8 dígitos."

        else:

            guardar_datos(
                dni,
                region,
                provincia,
                distrito,
                direccion
            )

            mensaje = "✅ Información registrada correctamente."

    return render_template_string(
        HTML,
        mensaje=mensaje
    )


@app.route("/excel")
def descargar_excel():

    crear_excel()

    return send_file(
        ARCHIVO,
        as_attachment=True
    )


if __name__ == "__main__":

    crear_excel()

    app.run(
        host="0.0.0.0",
        port=5000
    )